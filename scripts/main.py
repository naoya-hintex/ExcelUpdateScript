from config import load_config
from logger import setup_logger
from common import create_filePath
from pathlib import Path
import logging
import pandas as pd
import openpyxl as xl

# ＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
# 関数名           load_csv                                                             
# 処理概要         CSVファイルを読み込み、データフレーム型のデータで返却します。                                                          
# 作成者           naoya-hintex                                                           
# 引数             input_config: 設定情報                                                           
# 戻り値           df: CSVデータ                                                       
# ＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
def load_csv(input_config):

    # CSVファイルのパスを作成
    csv_path = create_filePath(input_config["csv_folder"], input_config["csv_file"])
    
    # CSVファイルを読み込み、データフレーム型で返却
    df = pd.read_csv(csv_path, encoding=input_config["encoding"])
    return df

# ＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
# 関数名           aggregate_daily                                                             
# 処理概要         CSVデータを集約しアウトプットデータを作成                                                          
# 作成者           naoya-hintex                                                           
# 引数             csv_data: CSVデータ                                                           
# 戻り値           summarycsv_data: アウトプットデータ                                                         
# ＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
def aggregate_daily(csv_data):

    # CSVデータを「date」で集約する
    # ※CSVには date / category / count / avg_process_time 列が存在する前提
    csv_data = csv_data.pivot_table(index='date', columns='category', values=['count', 'avg_process_time'])
    csv_data = csv_data[['count', 'avg_process_time']]
    return csv_data

# ＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
# 関数名           update_excel                                                             
# 処理概要         Excelファイルを更新する                                                          
# 作成者           naoya-hintex                                                           
# 引数             output_data: 集計後のデータ                                                           
# 引数             output_config: 設定情報                                                           
# ＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
def update_excel(output_data, output_config):

    # Excelファイルのパスを作成
    file_path = create_filePath(output_config["output_folder"], output_config["excel_file"])

    # シート名を取得
    sheet_name = output_config["sheet_name"]

    # Excelファイルをopenpyxlで読み込み
    book = xl.load_workbook(file_path)
    ws = book[sheet_name]

    # 「データ」シートにデータが存在する場合、2行目から最終行まで削除
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)

    # アウトプットデータをExcelファイルに追記
    with pd.ExcelWriter(file_path, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:

        writer._book = book 

        writer._sheets = {ws.title: ws for ws in book.worksheets}
        
        # 2行目から書き込む
        output_data.to_excel(writer, sheet_name=sheet_name, header=False, startrow=output_config["start_row"])

        # 罫線のクリア
        for row in ws.iter_rows(min_row=output_config["start_row"] + 1, max_row=ws.max_row, 
                                min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.border = xl.styles.Border()

# ＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
# 関数名           main                                                             
# 処理概要         メイン処理                                                          
# 作成者           naoya-hintex                                                           
# ＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊＊
def main():

    # configの読み込み
    config_data = load_config()

    # loggerの初期設定
    setup_logger(config_data["log"])
    logging.info("CSV転記処理を開始します。")

    try:

        # CSVファイルを取得 
        csv_data = load_csv(config_data["input"])
        logging.info("CSVファイルからデータ取得が完了しました。")

        # アウトプットデータの作成
        output_data = aggregate_daily(csv_data)
        logging.info("アウトプットデータの作成が完了しました。")

        # エクセルファイルを更新
        update_excel(output_data, config_data["output"])
        logging.info("エクセルファイルの更新が完了しました。")

    except Exception as e:
        logging.exception("予期しないエラーが発生しました。")
        raise

if __name__ == "__main__":
    main()